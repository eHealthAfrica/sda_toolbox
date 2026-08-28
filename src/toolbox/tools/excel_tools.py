import pandas as pd
from pandas import ExcelWriter
from openpyxl.styles import Side, Border
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from toolbox.mlos import PresetAttributes


def add_borders(excel_writer: ExcelWriter, sheet_name: str, column_count: int, row_count: int):
    book = excel_writer.book
    sheet = excel_writer.sheets[sheet_name]

    thin_side = Side(border_style='thin', color='000000')
    cell_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    for row in sheet.iter_rows(
            min_row=1,
            max_row=row_count + 1,
            min_col=1,
            max_col=column_count,
    ):
        for cell in row:
            cell.border = cell_border


def add_data_validation_rules(excel_writer: ExcelWriter, preset_attributes: PresetAttributes, sheet_name: str, data: pd.DataFrame):
    book = excel_writer.book
    sheet = book[sheet_name]
    preset_map: dict[str, list] = preset_attributes.preset_attribute_map()

    if len(preset_map)== 0:
        return

    for config_key, values in preset_map.items():
        col_idx = data.columns.get_loc(config_key) + 1
        column_letter = get_column_letter(col_idx)
        str_values = ','.join(values)
        std_values = f'"{str_values}"'
        dv = DataValidation(
            type="list", formula1=std_values, allowBlank=True,
            showErrorMessage=True, errorTitle = 'Invalid Selection',
            error = f'Please choose a valid option from the dropdown list for {config_key}.',
            errorStyle = 'stop'
        )

        sheet.add_data_validation(dv)
        start_row, end_row = 2, len(data) + 1
        cell_range = f"{column_letter}{start_row}:{column_letter}{end_row}"
        dv.add(cell_range)
